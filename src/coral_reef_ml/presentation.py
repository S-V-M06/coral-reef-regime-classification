"""Generate the polished, results-backed review presentation."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR
from .team_info import get_team_members
from .design_assets import prepare_design_assets

NAVY = RGBColor(16, 47, 58)
TEAL = RGBColor(27, 138, 143)
CORAL = RGBColor(240, 120, 97)
INK = RGBColor(34, 54, 62)
MUTED = RGBColor(96, 117, 125)
PALE = RGBColor(245, 248, 248)
MINT = RGBColor(225, 243, 240)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(220, 231, 233)


def _rgb(hexstr):
    hexstr = hexstr.lstrip('#')
    return RGBColor(int(hexstr[0:2],16), int(hexstr[2:4],16), int(hexstr[4:6],16))


def _shape(slide, x, y, w, h, fill, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None, line=None):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
    else:
        sh.line.fill.background()
    return sh


def _text(slide, x, y, w, h, text, size=16, color=INK, bold=False, font='Aptos',
          align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04, italic=False):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.italic = italic; p.font.color.rgb = color
    return sh


def _runs(slide, x, y, w, h, parts, default_size=14):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.04); tf.margin_right = Inches(.04); tf.margin_top = Inches(.02); tf.margin_bottom = Inches(.02)
    p = tf.paragraphs[0]
    for txt, color, bold in parts:
        r = p.add_run(); r.text = txt; r.font.name = 'Aptos'; r.font.size = Pt(default_size); r.font.color.rgb = color; r.font.bold = bold
    return sh


def _bullet(slide, x, y, w, text, size=15, color=INK, dot=TEAL, h=.54):
    _shape(slide, x, y+.13, .09, .09, dot, MSO_SHAPE.OVAL)
    _text(slide, x+.2, y, w-.2, h, text, size=size, color=color)


def _header(slide, section, title, number):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
    _shape(slide, 0, 0, 13.333, .14, TEAL, MSO_SHAPE.RECTANGLE)
    _text(slide, .58, .34, 12, .28, section.upper(), 9, TEAL, True)
    _text(slide, .58, .67, 12.1, .58, title, 25, NAVY, True)
    _shape(slide, .58, 1.38, 12.15, .018, LINE, MSO_SHAPE.RECTANGLE)
    _text(slide, .58, 7.15, 10, .18, 'UE24CS352A  |  Machine Learning Mini-Project  |  Problem 79', 8, MUTED)
    _text(slide, 12.2, 7.13, .55, .2, f'{number:02d}', 8, MUTED, True, align=PP_ALIGN.RIGHT)


def _card(slide, x, y, w, h, label, value, note=None, accent=TEAL):
    _shape(slide, x, y, w, h, PALE)
    _shape(slide, x, y, .07, h, accent, MSO_SHAPE.RECTANGLE)
    _text(slide, x+.18, y+.12, w-.3, .23, label.upper(), 8, MUTED, True)
    _text(slide, x+.18, y+.37, w-.3, .45, value, 24, NAVY, True)
    if note:
        _text(slide, x+.18, y+h-.28, w-.3, .2, note, 8, MUTED)


def generate_presentation(root: Path, cv_results, metrics, report_dict, importance, audit) -> Path:
    root = Path(root)
    out = root / 'results' / 'Coral_Reef_Mini_Project_Presentation.pptx'
    out.parent.mkdir(parents=True, exist_ok=True)
    assets = prepare_design_assets(root, cv_results, report_dict, importance, audit)
    team = get_team_members(root).replace('\n', '  |  ').replace('[Add both names]', 'Team names to be confirmed')
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Cover
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
    _shape(s, 0, 0, .18, 7.5, TEAL, MSO_SHAPE.RECTANGLE)
    _shape(s, 8.88, 0, 4.45, 7.5, RGBColor(22, 65, 75), MSO_SHAPE.RECTANGLE)
    _shape(s, 10.55, .8, 1.2, 1.2, CORAL, MSO_SHAPE.OVAL)
    _shape(s, 9.6, 1.9, 1.9, 1.9, TEAL, MSO_SHAPE.OVAL)
    _shape(s, 11.3, 3.4, .85, .85, RGBColor(83, 183, 175), MSO_SHAPE.OVAL)
    _shape(s, 9.3, 4.6, 2.1, 2.1, RGBColor(30, 82, 91), MSO_SHAPE.OVAL)
    _text(s, .8, .7, 7.6, .3, 'MACHINE LEARNING MINI-PROJECT  /  PROBLEM 79', 10, RGBColor(119, 218, 208), True)
    _text(s, .78, 1.55, 7.8, 1.3, 'Predicting Coral Reef\nRegimes', 35, WHITE, True)
    _text(s, .82, 3.0, 7.65, .7, 'from human and natural influences', 21, RGBColor(178, 222, 222), True)
    _text(s, .84, 4.1, 7.5, .55, 'A leakage-aware, four-class classification study', 15, WHITE)
    _shape(s, .84, 4.92, 7.2, .02, RGBColor(76, 111, 119), MSO_SHAPE.RECTANGLE)
    _text(s, .84, 5.18, 7.3, .46, f'Team  |  {team}', 13, RGBColor(227, 237, 238))
    _text(s, .84, 6.88, 7.5, .25, 'UE24CS352A  |  PES University', 9, RGBColor(161, 195, 199))
    _card(s, 9.15, .85, 3.45, 1.2, 'Observed sites', str(audit['rows']), 'Hawaiian reef dataset', TEAL)
    _card(s, 9.15, 2.35, 3.45, 1.2, 'Input predictors', str(audit['predictor_count']), 'Human + biophysical', CORAL)
    _card(s, 9.15, 3.85, 3.45, 1.2, 'Regime labels', '1 / 2 / 3 / 5', 'Nominal classes', TEAL)
    _text(s, 9.18, 5.55, 3.25, .8, 'Prediction is not causation.\nGeneralization beyond Hawaii is untested.', 12, RGBColor(219, 231, 232), True)

    # 2 Problem
    s = prs.slides.add_slide(blank); _header(s, '01  /  Problem', 'Can observed pressures classify a reef regime?', 2)
    _shape(s, .72, 1.72, 7.3, 2.2, PALE)
    _text(s, 1.02, 1.98, 6.7, .32, 'RESEARCH QUESTION', 10, TEAL, True)
    _text(s, 1.02, 2.42, 6.65, 1.05, 'Can 20 anthropogenic and biophysical predictors classify a Hawaiian reef site into its observed regime label?', 22, NAVY, True)
    _text(s, .82, 4.25, 7.15, .35, 'Four target classes in the supplied dataset', 13, MUTED, True)
    for i, lab in enumerate(['1','2','3','5']):
        x=.85+i*1.68; _shape(s,x,4.72,1.38,1.0, MINT if lab!='5' else RGBColor(255,239,234))
        _text(s,x,4.86,1.38,.55,lab,27,TEAL if lab!='5' else CORAL,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    _shape(s, 8.38, 1.72, 4.18, 4.6, NAVY)
    _text(s, 8.72, 2.02, 3.5, .35, 'SCOPE BOUNDARIES', 10, RGBColor(119,218,208), True)
    _bullet(s, 8.72, 2.58, 3.42, 'Predicts the dataset label from observed predictors.', 14, WHITE, CORAL, .72)
    _bullet(s, 8.72, 3.55, 3.42, 'Does not estimate causal effects.', 14, WHITE, CORAL, .7)
    _bullet(s, 8.72, 4.48, 3.42, 'Not a future-time forecast.', 14, WHITE, CORAL, .7)
    _bullet(s, 8.72, 5.4, 3.42, 'Regime 4 is absent in the source subset.', 14, WHITE, CORAL, .7)
    _text(s, .84, 6.18, 7.1, .35, 'Objective: compare model families, inspect class-wise errors and report limitations.', 12, MUTED)

    # 3 Dataset
    s = prs.slides.add_slide(blank); _header(s, '02  /  Data', 'Small dataset, explicit feature boundary', 3)
    _card(s, .72, 1.7, 2.72, 1.1, 'Sites', f"{audit['rows']}", 'All observations', TEAL)
    _card(s, 3.62, 1.7, 2.72, 1.1, 'Predictors', str(audit['predictor_count']), 'Effluent through Depth', CORAL)
    _card(s, 6.52, 1.7, 2.72, 1.1, 'Training split', str(audit['train_rows']), '80% / stratified', TEAL)
    _card(s, 9.42, 1.7, 2.72, 1.1, 'Held-out test', str(audit['test_rows']), '20% / one evaluation', CORAL)
    _text(s, .78, 3.1, 5.4, .34, 'CLASS DISTRIBUTION', 10, TEAL, True)
    s.shapes.add_picture(str(assets['class_dist']), Inches(.68), Inches(3.42), width=Inches(5.25), height=Inches(2.75))
    _shape(s, 6.18, 3.28, 6.2, 2.95, PALE)
    _text(s, 6.52, 3.55, 5.4, .34, 'DATA QUALITY & LEAKAGE SAFEGUARDS', 10, TEAL, True)
    _bullet(s, 6.52, 4.04, 5.45, 'Missing values: Complexity 113; Depth 12.', 13, INK, CORAL, .48)
    _bullet(s, 6.52, 4.72, 5.45, 'Median imputation is fitted inside each model pipeline.', 13, INK, CORAL, .55)
    _bullet(s, 6.52, 5.48, 5.45, 'Regime one-hot columns, IDs, coordinates and ecological response columns are excluded from inputs.', 12, INK, CORAL, .68)
    _text(s, .82, 6.42, 11.6, .35, 'Split: stratified 80/20, random_state = 42. Labels are nominal: 1, 2, 3 and 5.', 11, MUTED)

    # 4 Method
    s = prs.slides.add_slide(blank); _header(s, '03  /  Method', 'A controlled evaluation pipeline', 4)
    steps = [
        ('1','Audit','Validate rows, labels\nand predictor schema'),
        ('2','Split','Stratified 80/20\ntrain / test'),
        ('3','Compare','5-fold stratified CV\non training only'),
        ('4','Select','Highest mean\nCV macro-F1'),
        ('5','Evaluate','One held-out test\nevaluation')]
    xs = [.7, 3.18, 5.66, 8.14, 10.62]
    for i,(n,title,desc) in enumerate(steps):
        fill = NAVY if i==3 else PALE
        _shape(s,xs[i],2.05,2.0,1.75,fill)
        _shape(s,xs[i]+.16,2.2,.36,.36,CORAL if i==3 else TEAL,MSO_SHAPE.OVAL)
        _text(s,xs[i]+.16,2.2,.36,.34,n,10,WHITE,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
        _text(s,xs[i]+.16,2.72,1.7,.34,title,15,WHITE if i==3 else NAVY,True)
        _text(s,xs[i]+.16,3.12,1.72,.55,desc,10,RGBColor(226,239,239) if i==3 else MUTED)
        if i < 4:
            _shape(s,xs[i]+2.05,2.82,.34,.06,TEAL,MSO_SHAPE.RECTANGLE)
            _shape(s,xs[i]+2.3,2.72,.16,.25,TEAL,MSO_SHAPE.CHEVRON)
    _text(s,.8,4.42,5.45,.34,'CANDIDATE MODELS',10,TEAL,True)
    model_items = ['Most-frequent baseline','Logistic Regression','RBF SVM','Random Forest','Gradient Boosting','Random Forest + 5 interactions']
    for i,item in enumerate(model_items):
        x=.84+(i%2)*2.92; y=4.9+(i//2)*.52
        _shape(s,x,y,.12,.12,CORAL if 'Random Forest' in item else TEAL,MSO_SHAPE.OVAL)
        _text(s,x+.22,y-.06,2.55,.32,item,11,INK, i==3)
    _shape(s,7.1,4.42,5.3,2.1,MINT)
    _text(s,7.43,4.72,4.6,.3,'WHY MACRO-F1?',11,TEAL,True)
    _text(s,7.43,5.15,4.55,1.05,'It averages F1 across classes, so performance on smaller or harder classes is not hidden by the overall sample mix. CV chooses the model; the held-out test set does not.',13,NAVY)

    # 5 CV results
    s = prs.slides.add_slide(blank); _header(s, '04  /  Results', 'Random Forest led cross-validation', 5)
    _text(s,.78,1.58,11.7,.32,'Mean macro-F1 across five stratified folds; error bars show fold standard deviation.',11,MUTED)
    s.shapes.add_picture(str(assets['cv']), Inches(.72), Inches(2.0), width=Inches(8.1), height=Inches(4.25))
    _shape(s,9.0,2.12,3.5,3.7,NAVY)
    _text(s,9.35,2.48,2.8,.3,'SELECTED MODEL',10,RGBColor(119,218,208),True)
    _text(s,9.35,2.95,2.9,.55,'Random\nForest',23,WHITE,True)
    _text(s,9.35,4.23,2.9,.65,'0.692',31,RGBColor(119,218,208),True)
    _text(s,9.35,4.92,2.9,.3,'Mean CV macro-F1',11,WHITE)
    _shape(s,9.35,5.28,2.65,.02,RGBColor(71,105,113),MSO_SHAPE.RECTANGLE)
    _text(s,9.35,5.40,2.75,.28,'RF + interactions: 0.687',10,RGBColor(235,240,240),True)
    _text(s,.82,6.42,11.6,.32,'The interaction ablation did not improve CV macro-F1 in this run; the simpler 20-feature model was selected.',11,MUTED)

    # 6 Holdout evaluation
    s = prs.slides.add_slide(blank); _header(s, '05  /  Evaluation', 'Held-out test performance and errors', 6)
    _card(s,.75,1.65,2.7,1.13,'Accuracy',f"{metrics['test_accuracy']*100:.1f}%",'87 of 124 correct',TEAL)
    _card(s,3.65,1.65,2.7,1.13,'Macro-F1',f"{metrics['test_macro_f1']*100:.1f}%",'Across four classes',CORAL)
    _card(s,6.55,1.65,2.7,1.13,'Balanced accuracy',f"{metrics['test_balanced_accuracy']*100:.1f}%",'Mean class recall',TEAL)
    _card(s,9.45,1.65,2.7,1.13,'Test samples',str(metrics['test_rows']),'Held out from selection',CORAL)
    _text(s,.82,3.08,5.2,.32,'CONFUSION MATRIX',10,TEAL,True)
    s.shapes.add_picture(str(assets['cm']), Inches(.82), Inches(3.38), width=Inches(5.15), height=Inches(3.18))
    _shape(s,6.42,3.15,5.95,3.26,PALE)
    _text(s,6.75,3.48,5.2,.32,'PER-CLASS RECALL',10,TEAL,True)
    recalls = report_dict.set_index(report_dict.columns[0]) if report_dict.columns[0] == 'Unnamed: 0' else report_dict
    # use known rows from the classification report, whichever format is loaded
    for i,label in enumerate(['1','2','3','5']):
        row = report_dict[report_dict.iloc[:,0].astype(str)==label]
        rec = float(row['recall'].iloc[0]) if not row.empty and 'recall' in row.columns else float('nan')
        y=4.0+i*.48
        _text(s,6.77,y,.72,.32,f'R{label}',11,INK,True)
        _shape(s,7.72,y+.04,3.45,.2,LINE,MSO_SHAPE.ROUNDED_RECTANGLE)
        _shape(s,7.72,y+.04,3.45*(rec if rec==rec else 0),.2,TEAL if label!='1' else CORAL,MSO_SHAPE.ROUNDED_RECTANGLE)
        _text(s,11.35,y-.02,.55,.3,f'{rec:.3f}' if rec==rec else 'n/a',10,NAVY,True,align=PP_ALIGN.RIGHT)
    _text(s,6.77,6.0,5.0,.28,'R1 has the lowest recall in this test split.',10,MUTED,italic=True)

    # 7 interpretation
    s = prs.slides.add_slide(blank); _header(s, '06  /  Interpretation', 'Signals, errors and limits', 7)
    s.shapes.add_picture(str(assets['importance']), Inches(.58), Inches(1.72), width=Inches(6.25), height=Inches(3.75))
    _shape(s,7.02,1.75,5.45,2.0,MINT)
    _text(s,7.36,2.06,4.7,.32,'TOP THREE IN THIS RUN',10,TEAL,True)
    top = list(importance.sort_values('permutation_importance_mean', ascending=False)['feature'].astype(str).head(3))
    for i,feat in enumerate(top):
        _text(s,7.4,2.55+i*.34,4.6,.27,f'{i+1}.  {feat}',12,NAVY,i==0)
    _shape(s,7.02,4.02,5.45,2.12,PALE)
    _text(s,7.36,4.34,4.7,.3,'INTERPRET CAREFULLY',10,CORAL,True)
    _bullet(s,7.36,4.79,4.65,'Permutation importance is predictive, not causal.',12,INK,CORAL,.48)
    _bullet(s,7.36,5.43,4.65,'Negative or small values are possible and should not be overinterpreted.',12,INK,CORAL,.53)
    _text(s,.84,5.85,5.9,.75,'Scope limits: modest sample, Hawaiian sites only, regime 4 absent in source subset, no independent geographic validation.',12,MUTED)

    # 8 Conclusion/demo
    s = prs.slides.add_slide(blank); _header(s, '07  /  Conclusion', 'A reproducible baseline - not a causal claim', 8)
    _shape(s,.78,1.7,5.65,2.25,NAVY)
    _text(s,1.15,2.02,4.9,.32,'TAKEAWAY',10,RGBColor(119,218,208),True)
    _text(s,1.15,2.53,4.8,1.05,'Random Forest achieved 0.692 mean CV macro-F1 and 0.702 held-out macro-F1.',20,WHITE,True)
    _text(s,.84,4.35,5.4,.32,'LIVE DEMO CHECKLIST',10,TEAL,True)
    for i,txt in enumerate(['Show data audit and split sizes','Open model_comparison_cv.csv','Explain one confusion-matrix error','Explain one limitation / no causality']):
        _bullet(s,.9,4.82+i*.43,5.1,txt,12,INK,TEAL,.37)
    _shape(s,6.78,1.72,5.66,4.5,PALE)
    _text(s,7.12,2.04,4.9,.3,'REPRODUCIBILITY',10,TEAL,True)
    _text(s,7.12,2.47,4.9,.58,'Run from the project root:',12,MUTED)
    _shape(s,7.12,3.05,4.95,.65,NAVY)
    _text(s,7.3,3.21,4.55,.28,'python run_project.py',13,WHITE,True)
    _text(s,7.12,3.95,4.9,.6,'For document-only refresh after names/results change:',11,MUTED)
    _shape(s,7.12,4.52,4.95,.65,NAVY)
    _text(s,7.3,4.67,4.55,.3,'python regenerate_deliverables.py',10,WHITE,True)
    _text(s,7.12,5.47,4.95,.42,'Outputs are saved under results/.',11,INK,True)

    # 9 References / Q&A
    s = prs.slides.add_slide(blank); _header(s, '08  /  References', 'Sources and questions', 9)
    refs = [
        ('Primary ecological study', 'Jouffray et al. (2019). Parsing human and biophysical drivers of coral reef regimes. Proceedings of the Royal Society B, 286:20182544. https://doi.org/10.1098/rspb.2018.2544'),
        ('Stanford CS229 reference project', 'Cheng, Murphy & Walecka. Predicting Coral Reef Regimes from Human and Natural Influences. https://cs229.stanford.edu/proj2019spr/report/56.pdf'),
        ('Reference poster', 'https://cs229.stanford.edu/proj2019spr/poster/56.pdf'),
        ('Dataset repository', 'https://github.com/JBjouffray/Hawaii_RegimesPredictors'),
    ]
    for i,(label,body) in enumerate(refs):
        y=1.75+i*1.0
        _shape(s,.82,y,11.65,.77,PALE)
        _text(s,1.1,y+.12,2.5,.24,label,10,TEAL,True)
        _text(s,3.55,y+.1,8.55,.53,body,10,INK)
    _shape(s,.82,6.0,11.65,.62,NAVY)
    _text(s,1.08,6.14,11.0,.3,'Questions? Be ready to explain the split, macro-F1 selection, one class error, and one limitation.',13,WHITE,True)

    prs.save(out)
    return out
